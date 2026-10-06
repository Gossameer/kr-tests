"""
Выгрузка результатов в Excel — так, чтобы файл можно было сразу распечатать.

Листы:
  «Результаты» — класс, ФИО, вариант, балл за каждое задание, итог и процент.
                 Сортировка: класс → ФИО.
  «По умениям» — умение × класс: процент выполнения с заливкой
                 (меньше 50 — красный, 50–79 — жёлтый, 80 и больше — зелёный).
  «Задания»    — тексты заданий по вариантам с ответами и статистикой.

Оформление у всех листов одно: границы у всех ячеек, жирная залитая шапка
с переносом, закреплённые шапка и первые столбцы, автофильтр, ширина по
содержимому, альбомная страница, по ширине — на один лист, шапка повторяется
на каждой печатной странице.

Попадают только СДАННЫЕ и НЕ аннулированные работы.
"""

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.formulas import plain_text

# Пороги заливки на листе «По умениям».
LEVEL_LOW = 50
LEVEL_HIGH = 80

FILL_LOW = PatternFill("solid", fgColor="F8CBCB")
FILL_MID = PatternFill("solid", fgColor="FFE9B0")
FILL_HIGH = PatternFill("solid", fgColor="CDEBD3")
FILL_HEADER = PatternFill("solid", fgColor="D9E1F2")
FILL_TOTAL = PatternFill("solid", fgColor="EEF1F6")

THIN = Side(style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT_NAME = "Calibri"


def percent_of(correct: int | None, total: int | None) -> int:
    if not correct or not total:
        return 0
    return round(correct * 100 / total)


def level_fill(percent: int) -> PatternFill:
    if percent < LEVEL_LOW:
        return FILL_LOW
    if percent < LEVEL_HIGH:
        return FILL_MID
    return FILL_HIGH


def class_sort_key(name: str) -> tuple:
    """«5А» раньше «10А»: сначала число в начале, потом буквы."""
    match = re.match(r"\s*(\d+)\s*(.*)", name)
    if match:
        return (0, int(match.group(1)), match.group(2).lower())
    return (1, 0, name.lower())


def _width_of(value) -> float:
    """Ширина текста в «знаках Excel»: самая длинная строка ячейки."""
    if value is None:
        return 0
    return max(len(line) for line in str(value).split("\n")) if str(value) else 0


def finish_sheet(sheet, *, freeze: str, min_width: float = 5, max_width: float = 45,
                 wrap_columns: tuple[int, ...] = (), header_height: float | None = None) -> None:
    """
    Общее оформление листа: границы, шапка, ширины, печать.

    freeze — ячейка, выше и левее которой всё закреплено («C2» — шапка и два столбца).
    wrap_columns — номера столбцов с длинным текстом: у них перенос по словам.
    """
    last_row, last_column = sheet.max_row, sheet.max_column

    for row in sheet.iter_rows(min_row=1, max_row=last_row, max_col=last_column):
        for cell in row:
            cell.border = BORDER
            if cell.row == 1:
                cell.font = Font(name=FONT_NAME, bold=True)
                cell.fill = FILL_HEADER
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                if cell.font.name != FONT_NAME or not cell.font.bold:
                    cell.font = Font(name=FONT_NAME, bold=cell.font.bold)
                wrap = cell.column in wrap_columns
                horizontal = cell.alignment.horizontal or ("left" if wrap else None)
                cell.alignment = Alignment(horizontal=horizontal, vertical="top" if wrap else "center",
                                           wrap_text=wrap)

    # Ширина по содержимому. У шапки считаем самое длинное СЛОВО: она переносится.
    for column in range(1, last_column + 1):
        header = str(sheet.cell(row=1, column=column).value or "")
        widest = max((len(word) for word in re.split(r"[\s·]+", header)), default=0)
        for row in range(2, last_row + 1):
            widest = max(widest, _width_of(sheet.cell(row=row, column=column).value))
        limit = max_width if column in wrap_columns else max(max_width, 0)
        sheet.column_dimensions[get_column_letter(column)].width = min(max(widest + 2.5, min_width), limit)

    if header_height is not None:
        sheet.row_dimensions[1].height = header_height

    sheet.freeze_panes = freeze
    if last_row >= 1:
        sheet.auto_filter.ref = f"A1:{get_column_letter(last_column)}{last_row}"

    # Печать: альбомная, по ширине на один лист (в высоту — сколько получится),
    # шапка на каждой странице.
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"
    sheet.print_options.horizontalCentered = True
    sheet.page_margins.left = sheet.page_margins.right = 0.4
    sheet.page_margins.top = sheet.page_margins.bottom = 0.5
    sheet.oddFooter.center.text = "Стр. &P из &N"


def build_results_xlsx(
    test: dict,
    skills: list[dict],
    attempts: list[dict],
    task_marks: dict[int, dict[int, bool]],
    skill_cells: dict[tuple[int, int], dict],
    task_rows: list[dict],
) -> bytes:
    """
    attempts    — сданные, не аннулированные попытки (любой порядок);
    task_marks  — {id попытки: {номер задания в варианте: верно?}};
    skill_cells — {(id попытки, id умения): {"correct", "total"}};
    task_rows   — задания для листа «Задания» (со статистикой).
    """
    attempts = sorted(
        attempts,
        key=lambda item: (class_sort_key(item["student_class"]), item["student_name"].lower(), item["id"]),
    )
    tasks_count = max((item["max_score"] or 0 for item in attempts), default=0)
    tasks_count = max([tasks_count] + [max(marks, default=0) for marks in task_marks.values()])

    workbook = Workbook()

    # ---------- Лист 1: результаты ----------
    sheet = workbook.active
    sheet.title = "Результаты"
    sheet.append(
        ["Класс", "ФИО", "Вариант"]
        + [f"№ {number}" for number in range(1, tasks_count + 1)]
        + ["Итог", "Из", "%"]
    )
    for attempt in attempts:
        marks = task_marks.get(attempt["id"], {})
        sheet.append(
            [attempt["student_class"], attempt["student_name"], attempt["variant_no"]]
            # Задание без ответа — 0: оно не выполнено.
            + [1 if marks.get(number) else 0 for number in range(1, tasks_count + 1)]
            + [
                attempt["score"] or 0,
                attempt["max_score"] or 0,
                percent_of(attempt["score"], attempt["max_score"]),
            ]
        )
    for row in sheet.iter_rows(min_row=2, min_col=3):
        for cell in row:
            cell.alignment = Alignment(horizontal="center")
    for row in sheet.iter_rows(min_row=2, min_col=4, max_col=3 + tasks_count):
        for cell in row:
            if cell.value == 0:
                cell.fill = FILL_LOW
    percent_column = 6 + tasks_count
    for row in sheet.iter_rows(min_row=2, min_col=percent_column, max_col=percent_column):
        for cell in row:
            cell.number_format = '0"%"'
            cell.font = Font(name=FONT_NAME, bold=True)
    for row in sheet.iter_rows(min_row=2, min_col=percent_column - 2, max_col=percent_column - 2):
        for cell in row:
            cell.font = Font(name=FONT_NAME, bold=True)
    finish_sheet(sheet, freeze="C2", min_width=5)

    # ---------- Лист 2: умение × класс ----------
    classes = sorted({attempt["student_class"] for attempt in attempts}, key=class_sort_key)
    by_class: dict[str, list[dict]] = {name: [] for name in classes}
    for attempt in attempts:
        by_class[attempt["student_class"]].append(attempt)

    def skill_percent(skill_id: int, group: list[dict]) -> int | None:
        correct = total = 0
        for attempt in group:
            cell = skill_cells.get((attempt["id"], skill_id))
            if cell:
                correct += cell["correct"]
                total += cell["total"]
        return percent_of(correct, total) if total else None

    skills_sheet = workbook.create_sheet("По умениям")
    skills_sheet.append(["№", "Умение"] + classes + ["Все классы"])
    for skill in skills:
        skills_sheet.append(
            [skill["position"], skill["title"]]
            + [skill_percent(skill["id"], by_class[name]) for name in classes]
            + [skill_percent(skill["id"], attempts)]
        )
    skills_sheet.append(
        ["", "Сдали работ"] + [len(by_class[name]) for name in classes] + [len(attempts)]
    )
    total_row = skills_sheet.max_row
    for row in skills_sheet.iter_rows(min_row=2, max_row=total_row - 1, min_col=3):
        for cell in row:
            cell.alignment = Alignment(horizontal="center")
            if isinstance(cell.value, int):
                cell.fill = level_fill(cell.value)
                cell.number_format = '0"%"'
            else:
                cell.value = "—"
    for cell in skills_sheet[total_row]:
        cell.font = Font(name=FONT_NAME, bold=True)
        cell.fill = FILL_TOTAL
        if cell.column >= 3:
            cell.alignment = Alignment(horizontal="center")
    for row in skills_sheet.iter_rows(min_row=2, max_row=total_row, min_col=1, max_col=1):
        for cell in row:
            cell.alignment = Alignment(horizontal="center")
    finish_sheet(skills_sheet, freeze="C2", min_width=8, max_width=60, wrap_columns=(2,))

    # ---------- Лист 3: задания ----------
    tasks_sheet = workbook.create_sheet("Задания")
    tasks_sheet.append(
        ["Вариант", "№", "Умение", "Задание", "Формат", "Правильный ответ", "Решение", "Верно", "Ответов", "%"]
    )
    for task in task_rows:
        if task["answer_format"] == "choice":
            correct_answer = task["right_option_text"] or ""
            format_name = "выбор"
        else:
            correct_answer = " / ".join(task["accepted_answers"])
            format_name = "ввод"
        tasks_sheet.append(
            [
                task["variant_no"],
                task["position"],
                task["skill_title"],
                # В ячейке формулу не нарисовать — пишем её строкой: 3/5, x^2.
                plain_text(task["text"]),
                format_name,
                plain_text(correct_answer),
                plain_text(task["solution"]),
                task["correct"],
                task["answered"],
                percent_of(task["correct"], task["answered"]) if task["answered"] else None,
            ]
        )
    for row in tasks_sheet.iter_rows(min_row=2):
        for cell in row:
            if cell.column in (1, 2, 5, 8, 9, 10):
                cell.alignment = Alignment(horizontal="center")
            if cell.column == 10 and isinstance(cell.value, int):
                cell.number_format = '0"%"'
                cell.fill = level_fill(cell.value)
    finish_sheet(tasks_sheet, freeze="C2", max_width=50, wrap_columns=(3, 4, 6, 7))
    for letter, width in (("C", 26), ("D", 50), ("F", 22), ("G", 36)):
        tasks_sheet.column_dimensions[letter].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
