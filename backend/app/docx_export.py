"""
Файлы для печати: «Варианты» и «Ключ ответов» в формате Word (.docx).

«Варианты» — то, что раздают ученикам: каждый вариант с новой страницы, шапка
(название, «Вариант N», строка «ФИО ____ Класс ____ Дата ____»), задания
пронумерованы, варианты ответа — а) б) в) г), для ввода — строка для ответа.

«Ключ ответов» — для учителя: таблица «задание × вариант» с ответом и умением.

Формулы ($...$ в текстах заданий) вставляются НАСТОЯЩИМИ формулами Word (OMML):
их можно выделить и поправить в Word, они печатаются чётко при любом масштабе.
Формулу строим прямо из дерева разбора (app/formulas.py), без промежуточных
программ — на сервере не нужен ни Word, ни LaTeX. Если формулу не разобрать
(редкая команда), она печатается обычным текстом без знаков $: «x^2», а не «$x^2$».
"""

import io
import zipfile
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from app.formulas import FUNCTIONS, FormulaError, parse, split_formulas

FONT = "Times New Roman"
BODY_PT = 13
TITLE_PT = 14
LETTERS = "абвгдежзик"
# Сколько вариантов помещается в одну таблицу ключа на альбомном листе.
KEY_VARIANTS_PER_TABLE = 6

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


# =====================================================================
# Формула → OMML (формат формул Word)
# =====================================================================


def _is_cyrillic(text: str) -> bool:
    return any("а" <= char.lower() <= "я" or char.lower() == "ё" for char in text)


def _run(text: str, upright: bool, size_pt: float) -> str:
    """Кусок текста внутри формулы. upright — прямым шрифтом (слова, функции)."""
    style = '<m:rPr><m:sty m:val="p"/></m:rPr>' if upright else ""
    return (
        f"<m:r>{style}"
        f'<w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/>'
        f'<w:sz w:val="{int(size_pt * 2)}"/><w:szCs w:val="{int(size_pt * 2)}"/></w:rPr>'
        f'<m:t xml:space="preserve">{escape(text)}</m:t></m:r>'
    )


def _omml(node: tuple, size: float) -> str:
    kind = node[0]

    if kind == "row":
        parts: list[str] = []
        for child in node[1]:
            parts.append(_omml(child, size))
        return "".join(parts)

    if kind == "text":
        text = node[1]
        if text in FUNCTIONS:
            # «sin x»: без пробела Word склеил бы название функции с аргументом.
            return _run(text + " ", True, size)
        # Русские слова в формуле («см», «кг») — прямым шрифтом, как в учебнике.
        upright = _is_cyrillic(text) or (len(text) > 1 and text.strip().isalpha())
        return _run(text, upright, size)

    if kind == "frac":
        return f"<m:f><m:num>{_omml(node[1], size)}</m:num><m:den>{_omml(node[2], size)}</m:den></m:f>"

    if kind == "sqrt":
        if node[1] is None:
            return (
                '<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/>'
                f"<m:e>{_omml(node[2], size)}</m:e></m:rad>"
            )
        return f"<m:rad><m:deg>{_omml(node[1], size)}</m:deg><m:e>{_omml(node[2], size)}</m:e></m:rad>"

    if kind == "script":
        base, sub, sup = node[1], node[2], node[3]
        # «90^\circ» — это градус: обычный знак после числа, а не степень.
        if sup is not None and sup == ("text", "∘", True) and sub is None:
            return _omml(base, size) + _run("°", True, size)
        body = f"<m:e>{_omml(base, size)}</m:e>"
        if sub is not None and sup is not None:
            return (
                f"<m:sSubSup>{body}<m:sub>{_omml(sub, size)}</m:sub>"
                f"<m:sup>{_omml(sup, size)}</m:sup></m:sSubSup>"
            )
        if sup is not None:
            return f"<m:sSup>{body}<m:sup>{_omml(sup, size)}</m:sup></m:sSup>"
        return f"<m:sSub>{body}<m:sub>{_omml(sub, size)}</m:sub></m:sSub>"

    if kind == "delim":
        return (
            f'<m:d><m:dPr><m:begChr m:val="{escape(node[1])}"/>'
            f'<m:endChr m:val="{escape(node[2])}"/></m:dPr>'
            f"<m:e>{_omml(node[3], size)}</m:e></m:d>"
        )

    if kind == "accent":
        if node[1] == "bar":
            return (
                '<m:bar><m:barPr><m:pos m:val="top"/></m:barPr>'
                f"<m:e>{_omml(node[2], size)}</m:e></m:bar>"
            )
        return (
            '<m:acc><m:accPr><m:chr m:val="⃗"/></m:accPr>'
            f"<m:e>{_omml(node[2], size)}</m:e></m:acc>"
        )

    if kind == "cases":
        rows = "".join(f"<m:e>{_omml(row, size)}</m:e>" for row in node[1])
        return (
            '<m:d><m:dPr><m:begChr m:val="{"/><m:endChr m:val=""/></m:dPr>'
            f"<m:e><m:eqArr>{rows}</m:eqArr></m:e></m:d>"
        )

    return ""


def formula_element(latex: str, size_pt: float = BODY_PT):
    """Формула (без знаков $) → элемент <m:oMath> для вставки в абзац. Не разобралась — None."""
    try:
        tree = parse(latex)
    except FormulaError:
        return None
    return parse_xml(f'<m:oMath xmlns:m="{M_NS}" xmlns:w="{W_NS}">{_omml(tree, size_pt)}</m:oMath>')


# =====================================================================
# Абзацы
# =====================================================================


def _style_run(run, size_pt: float, bold: bool = False) -> None:
    run.font.name = FONT
    # Без этого кириллица в Word набирается шрифтом «для восточных языков» по умолчанию.
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run._element.rPr.rFonts.set(qn("w:cs"), FONT)
    run.font.size = Pt(size_pt)
    run.bold = bold


def add_text(paragraph, text: str, size_pt: float = BODY_PT, bold: bool = False) -> None:
    """Дописывает в абзац текст с формулами: куски $...$ — формулами Word."""
    for kind, chunk in split_formulas(text):
        if kind == "text":
            # Перевод строки в задании — перенос внутри абзаца.
            for number, line in enumerate(chunk.split("\n")):
                run = paragraph.add_run(line)
                if number:
                    run._element.insert(0, parse_xml(f'<w:br xmlns:w="{W_NS}"/>'))
                _style_run(run, size_pt, bold)
            continue
        element = formula_element(chunk, size_pt)
        if element is None:
            # Формулу не разобрать — печатаем как написано, но без знаков $.
            _style_run(paragraph.add_run(chunk.strip()), size_pt, bold)
        else:
            paragraph._p.append(element)


def _paragraph(document_or_cell, text: str = "", *, size: float = BODY_PT, bold: bool = False,
               align=None, space_after: float = 4, indent_mm: float = 0, keep_next: bool = False):
    paragraph = document_or_cell.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.15
    if indent_mm:
        paragraph.paragraph_format.left_indent = Mm(indent_mm)
    if align is not None:
        paragraph.alignment = align
    if keep_next:
        # Не отрывать условие от вариантов ответа на границе страницы.
        paragraph.paragraph_format.keep_with_next = True
    if text:
        add_text(paragraph, text, size, bold)
    return paragraph


def _new_document(landscape: bool = False):
    document = Document()
    style = document.styles["Normal"]
    style.font.name = FONT
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(BODY_PT)

    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Mm(297), Mm(210)
    section.left_margin = section.right_margin = Mm(18)
    section.top_margin = section.bottom_margin = Mm(15)
    return document


def _save(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# =====================================================================
# «Варианты»
# =====================================================================


def build_variants(test: dict, variants: dict[int, list[dict]]) -> bytes:
    """
    test — {"title", "subject"}; variants — {номер варианта: [задания по порядку]},
    задание — {"text", "answer_format", "options": [тексты]}.
    """
    document = _new_document()

    for index, (variant_no, tasks) in enumerate(sorted(variants.items())):
        if index:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            # Абзац с разрывом страницы пуст — убираем его высоту, чтобы шапка
            # следующего варианта начиналась с самого верха листа.
            last = document.paragraphs[-1]
            last.paragraph_format.space_after = Pt(0)
            last.paragraph_format.line_spacing = Pt(1)

        _paragraph(document, test["title"], size=TITLE_PT, bold=True,
                   align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        _paragraph(document, f"Вариант {variant_no}", size=TITLE_PT, bold=True,
                   align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
        _paragraph(
            document,
            "ФИО _______________________________  Класс _______  Дата ____________",
            space_after=12,
        )

        for number, task in enumerate(tasks, start=1):
            is_choice = task["answer_format"] == "choice"
            paragraph = _paragraph(document, space_after=3, keep_next=True)
            _style_run(paragraph.add_run(f"{number}. "), BODY_PT, bold=True)
            add_text(paragraph, task["text"])

            if is_choice:
                for position, option in enumerate(task["options"]):
                    letter = LETTERS[position] if position < len(LETTERS) else str(position + 1)
                    last_option = position == len(task["options"]) - 1
                    line = _paragraph(document, space_after=10 if last_option else 2,
                                      indent_mm=8, keep_next=not last_option)
                    _style_run(line.add_run(f"{letter}) "), BODY_PT)
                    add_text(line, option)
            else:
                _paragraph(document, "Ответ: ______________________________",
                           space_after=10, indent_mm=8)

    return _save(document)


# =====================================================================
# «Ключ ответов»
# =====================================================================


def answer_text(task: dict) -> str:
    """Ответ для ключа: «б) 2,5» у выбора, «3/5 или 0,6» у ввода."""
    if task["answer_format"] == "choice":
        for position, option in enumerate(task["options"]):
            if position == task["correct"]:
                letter = LETTERS[position] if position < len(LETTERS) else str(position + 1)
                return f"{letter}) {option}"
        return "—"
    return " или ".join(task["accepted_answers"]) or "—"


def _cell(cell, text: str, *, bold: bool = False, center: bool = False, size: float = 11) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(paragraph, text, size, bold)


def _repeat_header(row) -> None:
    """Строка шапки повторяется на каждой странице, если таблица не влезла на одну."""
    properties = row._tr.get_or_add_trPr()
    properties.append(parse_xml(f'<w:tblHeader xmlns:w="{W_NS}"/>'))
    properties.append(parse_xml(f'<w:cantSplit xmlns:w="{W_NS}"/>'))


def build_key(test: dict, variants: dict[int, list[dict]]) -> bytes:
    """
    Ключ: строки — задания, столбцы — варианты; в клетке ответ. Слева номер и умение.
    Задание — как в build_variants, плюс "correct", "accepted_answers", "skill_title".
    """
    document = _new_document(landscape=True)
    _paragraph(document, f"Ключ ответов — {test['title']}", size=TITLE_PT, bold=True,
               align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _paragraph(document, "Только для учителя. В клетках — правильные ответы по вариантам.",
               size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

    numbers = sorted(variants)
    rows_count = max((len(tasks) for tasks in variants.values()), default=0)

    for start in range(0, len(numbers), KEY_VARIANTS_PER_TABLE):
        chunk = numbers[start : start + KEY_VARIANTS_PER_TABLE]
        table = document.add_table(rows=1 + rows_count, cols=2 + len(chunk))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Ширины: номер узкий, умение пошире, остальное поровну между вариантами.
        usable = 297 - 36
        widths = [10, 62] + [(usable - 72) / len(chunk)] * len(chunk)
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Mm(width)

        header = table.rows[0]
        _repeat_header(header)
        for cell, title in zip(header.cells, ["№", "Умение"] + [f"Вариант {n}" for n in chunk]):
            _cell(cell, title, bold=True, center=True)
            cell._tc.get_or_add_tcPr().append(
                parse_xml(f'<w:shd xmlns:w="{W_NS}" w:val="clear" w:color="auto" w:fill="E7EAF0"/>')
            )

        for position in range(rows_count):
            cells = table.rows[position + 1].cells
            _cell(cells[0], str(position + 1), center=True)
            # Умение берём из первого варианта, где есть такое задание: порядок
            # умений во всех вариантах один.
            skill = next(
                (variants[n][position]["skill_title"] for n in numbers if position < len(variants[n])),
                "",
            )
            _cell(cells[1], skill)
            for cell, variant_no in zip(cells[2:], chunk):
                tasks = variants[variant_no]
                _cell(cell, answer_text(tasks[position]) if position < len(tasks) else "—", center=True)

        if start + KEY_VARIANTS_PER_TABLE < len(numbers):
            _paragraph(document, space_after=10)

    return _save(document)


def build_zip(code: str, variants_docx: bytes, key_docx: bytes) -> bytes:
    """Оба файла одним архивом: одна кнопка — одно скачивание."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(f"Варианты-{code}.docx", variants_docx)
        archive.writestr(f"Ключ-ответов-{code}.docx", key_docx)
    return buffer.getvalue()
