"""
Этап 3: файлы для печати (Word) и выгрузка результатов (Excel).

Файлы скачиваются у настоящего приложения и разбираются теми же библиотеками,
которыми созданы: проверяем содержимое, формулы и настройки печати.
"""

import io
import unittest
import zipfile

from docx import Document
from lxml import etree
from openpyxl import load_workbook

from app import docx_export, xlsx_export
from tests.support import Client, add_user, start_app

M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

SKILLS = [
    {"title": "Складывать дроби", "tasks_per_variant": 1, "answer_format": "input"},
    {"title": "Сравнивать числа", "tasks_per_variant": 1, "answer_format": "choice"},
]


def body() -> dict:
    return {
        "title": "Дроби и корни",
        "subject": "алгебра",
        "classes": ["7А", "10А"],
        "variants_count": 2,
        "shuffle": False,
        "skills": SKILLS,
        "variants": [
            {
                "variant_no": number,
                "tasks": [
                    {
                        "skill_index": 1,
                        "text": rf"Вычислите $\frac{{{number}}}{{7}} + 2\frac{{1}}{{7}}$ и $\sqrt{{16}}$.",
                        "answer_format": "input",
                        "accepted_answers": ["3/5", "0,6"],
                        "solution": r"$\frac{3}{5}$",
                    },
                    {
                        "skill_index": 2,
                        "text": r"Что больше $x^{2}$ при $x = 3$? Битая: $\frac{1}{2$",
                        "answer_format": "choice",
                        "options": [r"$2{,}5$", r"$\frac{19}{2}$", "десять"],
                        "correct": 1,
                    },
                ],
            }
            for number in (1, 2)
        ],
    }


def setUpModule() -> None:
    start_app()
    add_user("Экспортова Ольга", "export@example.org")
    add_user("Чужой Учитель", "stranger@example.org")


class Exports(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.teacher = Client().login("export@example.org")
        cls.stranger = Client().login("stranger@example.org")
        student = Client()
        status, created = cls.teacher.post("/api/tests", body())
        assert status == 201, created
        cls.test_id = created["id"]
        codes = {link["class_name"]: link["code"] for link in created["class_links"]}

        def solve(class_name: str, name: str, first: bool, second: bool, annul: bool = False) -> None:
            code = codes[class_name]
            started = student.post(f"/api/public/tests/{code}/start", {"student_name": name})[1]
            typed = next(t for t in started["tasks"] if t["answer_format"] == "input")
            choice = next(t for t in started["tasks"] if t["answer_format"] == "choice")
            option = choice["options"][1 if second else 0]["id"]
            status, _ = student.post(
                f"/api/public/tests/{code}/attempts/{started['attempt_id']}/submit",
                {
                    "attempt_token": started["attempt_token"],
                    "choices": {choice["id"]: option},
                    "inputs": {typed["id"]: "3/5" if first else "1"},
                },
            )
            assert status == 200
            if annul:
                cls.teacher.post(f"/api/tests/{cls.test_id}/attempts/{started['attempt_id']}/annul")

        solve("10А", "Яковлев Ян", True, True)
        solve("10А", "Борисов Боря", False, False)
        solve("7А", "Петров Пётр", True, False)
        solve("7А", "Аннулированный Вася", False, False, annul=True)
        solve("7А", "Андреева Аня", True, True)
        # Начал, но не сдал — в выгрузку не идёт.
        student.post(f"/api/public/tests/{codes['7А']}/start", {"student_name": "Незакончивший Коля"})

    def download(self, path: str, client=None) -> bytes:
        status, data = (client or self.teacher).get(f"/api/tests/{self.test_id}/{path}", raw=True)
        self.assertEqual(status, 200, data[:200])
        return data

    # ------------------------------------------------------------- Word

    def test_variants_docx(self) -> None:
        document = Document(io.BytesIO(self.download("print/variants.docx")))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)

        self.assertEqual(text.count("Дроби и корни"), 2)
        self.assertIn("Вариант 1", text)
        self.assertIn("Вариант 2", text)
        self.assertEqual(text.count("ФИО ___"), 2)
        self.assertEqual(text.count("Ответ: ___"), 2)
        for mark in ("1. ", "2. ", "а) ", "б) ", "в) "):
            self.assertIn(mark, text)
        # Сырой разметки нет: ни долларов, ни команд (битая формула — текстом без $).
        self.assertNotIn("$", text)
        self.assertNotIn(r"\sqrt", text)
        self.assertIn(r"\frac{1}{2", text)

        xml = document.element.xml
        # В каждом варианте: по 2 формулы в заданиях и 2 в вариантах ответа.
        self.assertEqual(xml.count("<m:oMath>") + xml.count("<m:oMath "), 12)
        self.assertEqual(xml.count("<m:f>"), 2 * 3)
        self.assertEqual(xml.count("<m:rad>"), 2)
        self.assertEqual(xml.count("<m:sSup>"), 2)
        # Каждый вариант с новой страницы: один разрыв между двумя вариантами.
        self.assertEqual(xml.count('w:type="page"'), 1)
        # Шрифт и размер.
        fonts = {run.font.name for p in document.paragraphs for run in p.runs if run.text.strip()}
        sizes = {run.font.size.pt for p in document.paragraphs for run in p.runs if run.text.strip()}
        self.assertEqual(fonts, {"Times New Roman"})
        self.assertTrue(sizes <= {12, 13, 14}, sizes)
        self.assertEqual(document.sections[0].orientation, 0)  # книжная

    def test_key_docx(self) -> None:
        document = Document(io.BytesIO(self.download("print/key.docx")))
        self.assertIn("Ключ ответов — Дроби и корни", document.paragraphs[0].text)
        table = document.tables[0]
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        self.assertEqual(rows[0], ["№", "Умение", "Вариант 1", "Вариант 2"])
        self.assertEqual(rows[1][:3], ["1", "1. Складывать дроби", "3/5 или 0,6"])
        self.assertEqual(rows[2][:2], ["2", "2. Сравнивать числа"])
        self.assertTrue(rows[2][2].startswith("б) "))
        # Ответ-формула в ключе — тоже формулой Word.
        self.assertEqual(table.rows[2].cells[2]._tc.xml.count("<m:f>"), 1)
        self.assertIn("w:tblHeader", table.rows[0]._tr.xml)
        self.assertEqual(document.sections[0].orientation, 1)  # альбомная

    def test_zip_has_both_files(self) -> None:
        archive = zipfile.ZipFile(io.BytesIO(self.download("print/print.zip")))
        names = sorted(archive.namelist())
        self.assertEqual(len(names), 2)
        self.assertTrue(names[0].startswith("Варианты-") and names[1].startswith("Ключ-ответов-"))
        for name in names:
            Document(io.BytesIO(archive.read(name)))  # оба открываются

    def test_only_owner_downloads(self) -> None:
        for path in ("print/print.zip", "print/variants.docx", "print/key.docx", "export.xlsx"):
            status, _ = self.stranger.get(f"/api/tests/{self.test_id}/{path}", raw=True)
            self.assertEqual(status, 403, path)
        self.assertEqual(self.teacher.get(f"/api/tests/{self.test_id}/print/other.docx", raw=True)[0], 404)
        self.assertEqual(Client().get(f"/api/tests/{self.test_id}/print/print.zip", raw=True)[0], 401)

    # ------------------------------------------------------------- Excel

    def test_results_sheet(self) -> None:
        workbook = load_workbook(io.BytesIO(self.download("export.xlsx")))
        self.assertEqual(workbook.sheetnames, ["Результаты", "По умениям", "Задания"])
        sheet = workbook["Результаты"]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]

        self.assertEqual(rows[0], ["Класс", "ФИО", "Вариант", "№ 1", "№ 2", "Итог", "Из", "%"])
        # Класс → ФИО; 7А раньше 10А; аннулированной и несданной работ нет.
        self.assertEqual(
            [(row[0], row[1]) for row in rows[1:]],
            [("7А", "Андреева Аня"), ("7А", "Петров Пётр"), ("10А", "Борисов Боря"), ("10А", "Яковлев Ян")],
        )
        self.assertEqual(rows[1][3:], [1, 1, 2, 2, 100])
        self.assertEqual(rows[2][3:], [1, 0, 1, 2, 50])
        self.assertEqual(rows[3][3:], [0, 0, 0, 2, 0])

        for row in sheet.iter_rows():
            for cell in row:
                self.assertEqual(cell.border.left.style, "thin", cell.coordinate)
                self.assertEqual(cell.border.bottom.style, "thin", cell.coordinate)
        for cell in sheet[1]:
            self.assertTrue(cell.font.bold)
            self.assertTrue(cell.alignment.wrap_text)
            self.assertEqual(cell.fill.fgColor.rgb, "00D9E1F2")
        self.check_print_setup(sheet, "A1:H5")
        # Ширина по содержимому: ФИО шире номера задания.
        self.assertGreater(sheet.column_dimensions["B"].width, sheet.column_dimensions["D"].width + 5)

    def check_print_setup(self, sheet, filter_range: str) -> None:
        self.assertEqual(sheet.freeze_panes, "C2")  # шапка и столбцы «Класс», «ФИО»
        self.assertEqual(sheet.auto_filter.ref, filter_range)
        self.assertEqual(sheet.page_setup.orientation, "landscape")
        self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual((sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight), (1, 0))
        self.assertEqual(sheet.print_title_rows, "$1:$1")

    def test_skills_sheet(self) -> None:
        sheet = load_workbook(io.BytesIO(self.download("export.xlsx")))["По умениям"]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        self.assertEqual(rows[0], ["№", "Умение", "7А", "10А", "Все классы"])
        self.assertEqual(rows[1], [1, "Складывать дроби", 100, 50, 75])
        self.assertEqual(rows[2], [2, "Сравнивать числа", 50, 50, 50])
        self.assertEqual(rows[3], [None, "Сдали работ", 2, 2, 4])

        red, yellow, green = "00F8CBCB", "00FFE9B0", "00CDEBD3"
        self.assertEqual(sheet["C2"].fill.fgColor.rgb, green)   # 100
        self.assertEqual(sheet["D2"].fill.fgColor.rgb, yellow)  # 50
        self.assertEqual(sheet["E2"].fill.fgColor.rgb, yellow)  # 75
        self.assertEqual(sheet["C2"].number_format, '0"%"')
        self.check_print_setup(sheet, "A1:E4")
        self.assertEqual(xlsx_export.level_fill(49).fgColor.rgb, red)
        self.assertEqual(xlsx_export.level_fill(79).fgColor.rgb, yellow)
        self.assertEqual(xlsx_export.level_fill(80).fgColor.rgb, green)

    def test_tasks_sheet_has_no_raw_markup(self) -> None:
        sheet = load_workbook(io.BytesIO(self.download("export.xlsx")))["Задания"]
        texts = [str(cell) for row in sheet.iter_rows(min_row=2, values_only=True) for cell in row if cell]
        self.assertIn("Вычислите 1/7+2 1/7 и √16.", texts)
        self.assertIn("19/2", texts)
        # Ответы аннулированной попытки в статистику заданий не попали.
        answered = [row[8] for row in sheet.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(sum(answered), 8)


class Formulas(unittest.TestCase):
    def omml(self, latex: str) -> str:
        element = docx_export.formula_element(latex)
        self.assertIsNotNone(element, latex)
        return etree.tostring(element, encoding="unicode")

    def test_constructs(self) -> None:
        self.assertIn("<m:f>", self.omml(r"\frac{a+b}{2}"))
        self.assertIn("<m:deg>", self.omml(r"\sqrt[3]{x}"))
        self.assertIn('m:degHide m:val="1"', self.omml(r"\sqrt{x}"))
        self.assertIn("<m:sSubSup>", self.omml(r"a_1^2"))
        self.assertIn("<m:sSub>", self.omml(r"a_n"))
        self.assertIn('m:begChr m:val="("', self.omml(r"\left(\frac{1}{2}\right)^2"))
        self.assertIn("<m:eqArr>", self.omml(r"\begin{cases} x+y=5 \\ x-y=1 \end{cases}"))
        self.assertIn("<m:bar>", self.omml(r"\overline{AB}"))
        self.assertIn("<m:acc>", self.omml(r"\vec{a}"))
        degree = self.omml(r"90^\circ")
        self.assertIn("°", degree)
        self.assertNotIn("<m:sSup>", degree)
        # Знаки — настоящими символами, не командами.
        signs = self.omml(r"3 \cdot 4 \le x \ne \pi")
        for char in "·≤≠π":
            self.assertIn(char, signs)
        # Русские единицы — прямым шрифтом.
        self.assertIn('<m:sty m:val="p"/>', self.omml(r"5 \text{ см}"))

    def test_unparsed_formula_falls_back_to_text(self) -> None:
        self.assertIsNone(docx_export.formula_element(r"\unknown{x}"))
        document = Document(
            io.BytesIO(
                docx_export.build_variants(
                    {"title": "Т"},
                    {1: [{"text": r"Формула $\unknown{x}$ и $x<y&z$", "answer_format": "input", "options": []}]},
                )
            )
        )
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn(r"Формула \unknown{x} и x<y&z", text)
        self.assertNotIn("$", text)

    def test_many_variants_split_key_tables(self) -> None:
        task = {"text": "т", "answer_format": "input", "options": [], "correct": None,
                "accepted_answers": ["1"], "skill_title": "1. У"}
        document = Document(io.BytesIO(docx_export.build_key({"title": "Т"}, {n: [task] for n in range(1, 15)})))
        self.assertEqual([len(table.columns) for table in document.tables], [8, 8, 4])


if __name__ == "__main__":
    unittest.main()
