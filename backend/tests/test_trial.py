"""
Пробное прохождение: учитель или администратор открыл ссылку для учеников, будучи
вошедшим, и прошёл работу сам. Попытка видна с пометкой, но не попадает в баллы
класса, умения, статистику, «ошибки учеников» и Excel и не мешает настоящим ученикам.
"""

import io
import unittest

import psycopg
from openpyxl import load_workbook

from tests.support import Client, add_user, drop_database, migrations, scratch_database, start_app


def body() -> dict:
    return {
        "title": "Пробная работа",
        "subject": "алгебра",
        "classes": ["9А"],
        "variants_count": 2,
        "shuffle": False,
        "skills": [{"title": "Считать", "tasks_per_variant": 1, "answer_format": "input"}],
        "variants": [
            {"variant_no": n, "tasks": [{"skill_index": 1, "text": f"В{n}: $2 + 2$. Ответ запишите целым числом.",
                                         "answer_format": "input", "accepted_answers": ["4"]}]}
            for n in (1, 2)
        ],
    }


def setUpModule() -> None:
    start_app()
    add_user("Пробова Полина", "trial@example.org")
    add_user("Чужой Учитель", "trial-other@example.org")


class Trial(unittest.TestCase):
    def setUp(self) -> None:
        self.teacher = Client().login("trial@example.org")
        status, created = self.teacher.post("/api/tests", body())
        self.assertEqual(status, 201)
        self.test_id = created["id"]
        self.code = created["class_links"][0]["code"]

    def solve(self, client: Client, name: str, answer: str, device: str = "") -> dict:
        status, started = client.post(f"/api/public/tests/{self.code}/start",
                                      {"student_name": name, "device_id": device})
        self.assertEqual(status, 201, started)
        status, result = client.post(
            f"/api/public/tests/{self.code}/attempts/{started['attempt_id']}/submit",
            {"attempt_token": started["attempt_token"], "choices": {}, "inputs": {started["tasks"][0]["id"]: answer}},
        )
        self.assertEqual(status, 200, result)
        return started

    def test_trial_is_visible_but_not_counted(self) -> None:
        # Учитель «прошёл как ученик», ответив неверно, — на том же устройстве, что и класс.
        trial = self.solve(self.teacher, "Иванов Иван", "5", device="pc-1")
        self.assertTrue(trial["is_trial"])

        # Настоящий Иванов Иван с того же компьютера — не мешает ни ФИО, ни устройство,
        # и вариант ему достаётся первый: проба очередь не сдвинула.
        real = self.solve(Client(), "Иванов Иван", "4", device="pc-1")
        self.assertFalse(real["is_trial"])
        self.assertEqual(real["variant_no"], 1)

        results = self.teacher.get(f"/api/tests/{self.test_id}/results")[1]
        rows = {(row["student_name"], row["trial"]): row for row in results["attempts"]}
        self.assertEqual(set(rows), {("Иванов Иван", True), ("Иванов Иван", False)})
        self.assertEqual(results["attempts_count"], 1)
        self.assertEqual(results["class_links"][0]["attempts_count"], 1)
        self.assertEqual(results["skill_stats"][0]["percent"], 100)  # ошибка пробы не считается

        stats = self.teacher.get(f"/api/stats?test_id={self.test_id}")[1]
        self.assertEqual(stats["by_class"][0]["attempts_count"], 1)
        self.assertEqual(stats["by_class"][0]["average_percent"], 100)
        self.assertEqual({row["percent"] for row in stats["weak_skills"]}, {100})

        review = self.teacher.get(f"/api/tests/{self.test_id}/tasks-review")[1]["tasks"]
        self.assertEqual(sum(task["wrong"] for task in review), 0)

        content = self.teacher.get(f"/api/tests/{self.test_id}/export.xlsx", raw=True)[1]
        names = [row[1] for row in load_workbook(io.BytesIO(content))["Результаты"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(names, ["Иванов Иван"])
        mine = next(row for row in self.teacher.get("/api/my/tests")[1] if row["id"] == self.test_id)
        self.assertEqual(mine["attempts_count"], 1)

        # Удалить пробные прохождения: чужой не может, автор — да.
        self.assertEqual(Client().login("trial-other@example.org").request(
            "DELETE", f"/api/tests/{self.test_id}/trial-attempts")[0], 403)
        status, data = self.teacher.request("DELETE", f"/api/tests/{self.test_id}/trial-attempts")
        self.assertEqual((status, data), (200, {"deleted": 1}))
        results = self.teacher.get(f"/api/tests/{self.test_id}/results")[1]
        self.assertEqual([row["trial"] for row in results["attempts"]], [False])

    def test_trial_only_work_can_still_be_edited(self) -> None:
        self.solve(self.teacher, "Проба Учителя", "4")
        draft = self.teacher.get(f"/api/tests/{self.test_id}/draft")[1]
        self.assertFalse(draft["has_attempts"])
        status, _ = self.teacher.request("PUT", f"/api/tests/{self.test_id}", body())
        self.assertEqual(status, 200)
        self.assertEqual(self.teacher.get(f"/api/tests/{self.test_id}/results")[1]["attempts"], [])

    def test_student_without_login_is_not_trial(self) -> None:
        self.assertFalse(self.solve(Client(), "Обычный Ученик", "4")["is_trial"])


class Migration014(unittest.TestCase):
    NAME = "kr_tests_tmp_migration014"

    def test_up_down(self) -> None:
        url = scratch_database(self.NAME)
        self.addCleanup(drop_database, self.NAME)
        up = next(path for path in migrations() if path.name.startswith("014"))
        down = up.with_name(up.name.replace(".sql", ".down.sql"))
        with psycopg.connect(url, autocommit=True) as conn:
            for path in migrations(up_to="013"):
                conn.execute(path.read_text(encoding="utf-8"))
            teacher = conn.execute(
                "INSERT INTO users (full_name, email, password_hash) VALUES ('У', 'u@e.org', 'x') RETURNING id"
            ).fetchone()[0]
            test_id = conn.execute(
                "INSERT INTO tests (title, teacher_id, teacher_name, share_token) VALUES ('Т', %s, 'У', 'abcd1234') RETURNING id",
                (teacher,),
            ).fetchone()[0]
            conn.execute(up.read_text(encoding="utf-8"))
            conn.execute(up.read_text(encoding="utf-8"))
            for trial in (True, False):
                conn.execute(
                    "INSERT INTO attempts (test_id, student_name, name_key, attempt_token, is_trial) "
                    "VALUES (%s, 'Иванов Иван', 'иванов иван', %s, %s)",
                    (test_id, f"t{trial}", trial),
                )
            conn.execute(down.read_text(encoding="utf-8"))
            self.assertEqual(conn.execute("SELECT count(*) FROM attempts").fetchone()[0], 1)
            conn.execute(up.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
