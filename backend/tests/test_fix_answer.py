"""
«Исправить правильный ответ»: замена ответа (а не добавление), пересчёт сданных
работ, баллов, умений, статистики, Excel и ключа ответов; отмена; журнал.
Примеры — настоящие ошибки генерации из тестовой работы по логарифмам.
"""

import io
import unittest

import psycopg
from docx import Document
from openpyxl import load_workbook

from tests.support import Client, add_user, drop_database, migrations, query, scratch_database, start_app


def body() -> dict:
    return {
        "title": "Логарифмы (ошибки ИИ)",
        "subject": "алгебра",
        "classes": ["10А"],
        "variants_count": 1,
        "shuffle": False,
        "skills": [
            {"title": "Решать уравнения", "tasks_per_variant": 1, "answer_format": "input"},
            {"title": "Сравнивать логарифмы", "tasks_per_variant": 1, "answer_format": "choice"},
        ],
        "variants": [{"variant_no": 1, "tasks": [
            {"skill_index": 1, "answer_format": "input", "accepted_answers": ["16"],
             "text": "Найдите произведение корней уравнения $\\log_4^2 x - 3\\log_4 x + 2 = 0$. Ответ запишите целым числом."},
            {"skill_index": 2, "answer_format": "choice", "correct": 3,
             "text": "Какое выражение наибольшее?",
             "options": ["$\\log_2 8$", "$\\log_3 9$", "$\\log_5 25$", "$\\log_{10} 100$"]},
        ]}],
    }


def setUpModule() -> None:
    start_app()
    add_user("Исправлялова Ирина", "fixanswer@example.org")


class FixCorrectAnswer(unittest.TestCase):
    def setUp(self) -> None:
        self.teacher = Client().login("fixanswer@example.org")
        status, created = self.teacher.post("/api/tests", body())
        self.assertEqual(status, 201, created)
        self.test_id = created["id"]
        code = created["class_links"][0]["code"]
        student = Client()
        self.ids = {}
        # Петров и Сидоров решили верно (64, log_2 8), Иванов согласился с ошибкой ИИ.
        for name, typed, option in (("Петров Пётр", "64", 0), ("Сидоров Семён", "64", 0), ("Иванов Иван", "16", 3)):
            started = student.post(f"/api/public/tests/{code}/start", {"student_name": name})[1]
            typed_task = next(t for t in started["tasks"] if t["answer_format"] == "input")
            choice = next(t for t in started["tasks"] if t["answer_format"] == "choice")
            self.ids = {"input": typed_task["id"], "choice": choice["id"],
                        "options": [o["id"] for o in choice["options"]], **self.ids}
            student.post(f"/api/public/tests/{code}/attempts/{started['attempt_id']}/submit", {
                "attempt_token": started["attempt_token"],
                "choices": {choice["id"]: choice["options"][option]["id"]},
                "inputs": {typed_task["id"]: typed},
            })
            self.ids[name] = started["attempt_id"]

    def review(self, task: str, action: str, **extra):
        return self.teacher.post(f"/api/tests/{self.test_id}/tasks/{self.ids[task]}/review", {"action": action, **extra})

    def scores(self) -> dict[str, tuple[int, int]]:
        data = self.teacher.get(f"/api/tests/{self.test_id}/results")[1]
        return {row["student_name"]: (row["score"], row["percent"]) for row in data["attempts"]}

    def excel_rows(self) -> dict[str, tuple]:
        content = self.teacher.get(f"/api/tests/{self.test_id}/export.xlsx", raw=True)[1]
        sheet = load_workbook(io.BytesIO(content))["Результаты"]
        return {row[1]: row[3:] for row in sheet.iter_rows(min_row=2, values_only=True)}

    def key_text(self) -> str:
        content = self.teacher.get(f"/api/tests/{self.test_id}/print/key.docx", raw=True)[1]
        return "\n".join(cell.text for row in Document(io.BytesIO(content)).tables[0].rows for cell in row.cells)

    def test_input_answer_replaced_everything_rescored_and_undone(self) -> None:
        self.assertEqual(self.scores(), {"Петров Пётр": (0, 0), "Сидоров Семён": (0, 0), "Иванов Иван": (2, 100)})

        # Сводка «ошибки учеников»: кто как ответил — фамилии короткие.
        tasks = self.teacher.get(f"/api/tests/{self.test_id}/tasks-review")[1]["tasks"]
        summary = next(task for task in tasks if task["task_id"] == self.ids["input"])
        self.assertEqual(summary["top_wrong"][0]["answer"], "64")
        self.assertEqual(summary["top_wrong"][0]["students"], ["Петров П.", "Сидоров С."])
        self.assertEqual(summary["top_wrong"][0]["more"], 0)
        self.assertIsNone(summary["original_answer"])

        status, data = self.review("input", "set_answer", answer="64")
        self.assertEqual(status, 200, data)
        self.assertEqual((data["accepted_answers"], data["original_answers"]), (["64"], ["16"]))
        scores = self.scores()
        self.assertEqual(scores["Петров Пётр"], (1, 50))
        self.assertEqual(scores["Иванов Иван"], (1, 50))  # 16 теперь неверно — ответ заменён, а не добавлен

        # Разбор, умения, статистика, Excel и ключ ответов — по новому ответу.
        detail = self.teacher.get(f"/api/tests/{self.test_id}/attempts/{self.ids['Петров Пётр']}")[1]
        item = next(i for i in detail["items"] if i["task_id"] == self.ids["input"])
        self.assertTrue(item["is_correct"])
        self.assertEqual((item["correct_answer"], item["original_answer"]), ("64", "16"))
        results = self.teacher.get(f"/api/tests/{self.test_id}/results")[1]
        skill = next(stat for stat in results["skill_stats"] if stat["position"] == 1)
        self.assertEqual((skill["correct"], skill["total"]), (2, 3))
        stats = self.teacher.get(f"/api/stats?test_id={self.test_id}")[1]
        self.assertEqual(stats["by_class"][0]["average_percent"], 50)
        self.assertEqual(self.excel_rows()["Петров Пётр"][:2], (1, 0))
        self.assertIn("64", self.key_text())
        self.assertNotIn("16", self.key_text().split("Вариант 1")[1])

        # Несколько записей через «|»; повторная правка не теряет исходный ответ.
        self.assertEqual(self.review("input", "set_answer", answer="64 | шестьдесят четыре")[0], 200)
        self.assertEqual(query("SELECT original_answers FROM tasks WHERE id = %s", (self.ids["input"],))[0][0], ["16"])

        # Отмена — всё как было.
        status, data = self.review("input", "restore_answer")
        self.assertEqual((status, data["accepted_answers"], data["original_answers"]), (200, ["16"], None))
        self.assertEqual(self.scores()["Иванов Иван"], (2, 100))
        self.assertEqual(self.review("input", "restore_answer")[0], 409)

        log = self.teacher.get(f"/api/tests/{self.test_id}/review-log")[1]
        self.assertEqual([row["action"] for row in log], ["restore_answer", "set_answer", "set_answer"])
        self.assertIn("было «16», стало «64»", log[2]["detail"])
        self.assertEqual(log[2]["action_name"], "исправил правильный ответ")

    def test_choice_other_option_marked_correct(self) -> None:
        log2_8 = self.ids["options"][0]
        status, _ = self.review("choice", "set_option", option_id=log2_8)
        self.assertEqual(status, 200)
        scores = self.scores()
        self.assertEqual(scores["Петров Пётр"], (1, 50))
        self.assertEqual(scores["Иванов Иван"], (1, 50))
        self.assertEqual(query("SELECT text FROM task_options WHERE task_id = %s AND is_correct",
                               (self.ids["choice"],))[0][0], "$\\log_2 8$")
        # В ключе ответов — новый верный вариант «а)» (формула внутри — формулой Word).
        self.assertIn("а) ", self.key_text())
        self.assertNotIn("г) ", self.key_text())

        tasks = self.teacher.get(f"/api/tests/{self.test_id}/tasks-review")[1]["tasks"]
        summary = next(task for task in tasks if task["task_id"] == self.ids["choice"])
        self.assertEqual(summary["original_answer"], "$\\log_{10} 100$")

        self.assertEqual(self.review("choice", "set_option", option_id=log2_8)[0], 400)  # уже верный
        self.assertEqual(self.review("choice", "set_answer", answer="1")[0], 400)
        self.assertEqual(self.review("input", "set_option", option_id=log2_8)[0], 400)
        self.assertEqual(self.review("input", "set_answer", answer=" | ")[0], 422)

        self.assertEqual(self.review("choice", "restore_answer")[0], 200)
        self.assertEqual(self.scores()["Иванов Иван"], (2, 100))


class Migration013(unittest.TestCase):
    NAME = "kr_tests_tmp_migration013"

    def test_up_down_restores_original_answers(self) -> None:
        url = scratch_database(self.NAME)
        self.addCleanup(drop_database, self.NAME)
        up = next(path for path in migrations() if path.name.startswith("013"))
        down = up.with_name(up.name.replace(".sql", ".down.sql"))
        with psycopg.connect(url, autocommit=True) as conn:
            for path in migrations(up_to="012"):
                conn.execute(path.read_text(encoding="utf-8"))
            teacher = conn.execute(
                "INSERT INTO users (full_name, email, password_hash) VALUES ('У', 'u@e.org', 'x') RETURNING id"
            ).fetchone()[0]
            test_id = conn.execute(
                "INSERT INTO tests (title, teacher_id, teacher_name, share_token) VALUES ('Т', %s, 'У', 'abcd1234') RETURNING id",
                (teacher,),
            ).fetchone()[0]
            skill = conn.execute("INSERT INTO skills (test_id, title) VALUES (%s, 'У') RETURNING id", (test_id,)).fetchone()[0]
            task = conn.execute(
                "INSERT INTO tasks (test_id, skill_id, variant_no, text, accepted_answers) VALUES (%s, %s, 1, 'т', '{16}') RETURNING id",
                (test_id, skill),
            ).fetchone()[0]
            conn.execute(up.read_text(encoding="utf-8"))
            conn.execute(up.read_text(encoding="utf-8"))
            conn.execute("UPDATE tasks SET original_answers = accepted_answers, accepted_answers = '{64}' WHERE id = %s", (task,))
            conn.execute(down.read_text(encoding="utf-8"))
            self.assertEqual(conn.execute("SELECT accepted_answers FROM tasks").fetchone()[0], ["16"])
            columns = {r[0] for r in conn.execute("SELECT column_name FROM information_schema.columns WHERE table_name='tasks'")}
            self.assertFalse(columns & {"original_answers", "original_correct_option"})
            conn.execute(up.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
