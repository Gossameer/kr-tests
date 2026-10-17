"""
Исправление проведённых работ: «добавить верный ответ», «засчитать всем»,
«исключить из оценки», их отмена, пересчёт результатов и журнал.

Пример из жизни: «Найдите длину промежутка области определения y = log9(5 − x)».
Область бесконечная, «правильный» ответ — «бесконечность», а ученица верно
написала промежуток «(-бесконечность;5)» — ей поставили ошибку.
"""

import io
import unittest

import psycopg
from openpyxl import load_workbook

from tests.support import Client, add_user, drop_database, migrations, query, scratch_database, start_app

LOG_TASK = "Найдите длину промежутка области определения функции $y = \\log_9(5 - x)$."


def body() -> dict:
    return {
        "title": "Логарифмы",
        "subject": "алгебра",
        "classes": ["10А"],
        "variants_count": 2,
        "shuffle": False,
        "skills": [
            {"title": "Находить область определения", "tasks_per_variant": 1, "answer_format": "input"},
            {"title": "Решать уравнения", "tasks_per_variant": 1, "answer_format": "choice"},
        ],
        "variants": [
            {
                "variant_no": number,
                "tasks": [
                    {
                        "skill_index": 1,
                        "text": LOG_TASK if number == 1 else "Сколько целых чисел в области определения?",
                        "answer_format": "input",
                        "accepted_answers": ["бесконечность", "беск"],
                    },
                    {
                        "skill_index": 2,
                        "text": f"Вариант {number}: корень уравнения $\\log_2 x = 3$",
                        "answer_format": "choice",
                        "options": ["8", "6", "9"],
                        "correct": 0,
                    },
                ],
            }
            for number in (1, 2)
        ],
    }


def setUpModule() -> None:
    start_app()
    add_user("Логарифмова Ольга", "review@example.org")
    add_user("Посторонний Учитель", "other-review@example.org")
    add_user("Главный Админ", "review-admin@example.org", "admin")


class Review(unittest.TestCase):
    def setUp(self) -> None:
        self.teacher = Client().login("review@example.org")
        self.student = Client()
        status, created = self.teacher.post("/api/tests", body())
        self.assertEqual(status, 201, created)
        self.test_id = created["id"]
        self.code = created["class_links"][0]["code"]
        # Вариант 1 (по очереди): Аня, Вера; вариант 2: Боря, Гоша.
        self.attempts = {}
        for name, typed, option in (
            ("Аня Первая", "(-бесконечность;5)", 0),
            ("Боря Второй", "5", 1),
            ("Вера Третья", "( - бесконечность ; 5 )", 1),
            ("Гоша Четвёртый", "", 0),
        ):
            self.attempts[name] = self.submit(name, typed, option)

    def submit(self, name: str, typed: str, option_index: int) -> dict:
        started = self.student.post(f"/api/public/tests/{self.code}/start", {"student_name": name})[1]
        typed_task = next(t for t in started["tasks"] if t["answer_format"] == "input")
        choice = next(t for t in started["tasks"] if t["answer_format"] == "choice")
        status, result = self.student.post(
            f"/api/public/tests/{self.code}/attempts/{started['attempt_id']}/submit",
            {
                "attempt_token": started["attempt_token"],
                "choices": {choice["id"]: choice["options"][option_index]["id"]},
                "inputs": {typed_task["id"]: typed} if typed else {},
            },
        )
        self.assertEqual(status, 200, result)
        return {**result, "variant_no": started["variant_no"], "input_id": typed_task["id"],
                "choice_id": choice["id"], "options": choice["options"]}

    def scores(self) -> dict[str, tuple[int, int, int]]:
        data = self.teacher.get(f"/api/tests/{self.test_id}/results")[1]
        return {row["student_name"]: (row["score"], row["max_score"], row["percent"]) for row in data["attempts"]}

    def review(self, task_id: int, action: str, client=None, **extra):
        return (client or self.teacher).post(
            f"/api/tests/{self.test_id}/tasks/{task_id}/review", {"action": action, **extra}
        )

    def task_summary(self, task_id: int) -> dict:
        tasks = self.teacher.get(f"/api/tests/{self.test_id}/tasks-review")[1]["tasks"]
        return next(task for task in tasks if task["task_id"] == task_id)

    def test_add_answer_rescores_everything_and_can_be_undone(self) -> None:
        task_id = self.attempts["Аня Первая"]["input_id"]
        self.assertEqual(self.scores()["Аня Первая"], (1, 2, 50))

        summary = self.task_summary(task_id)
        self.assertEqual((summary["answered"], summary["wrong"]), (2, 2))
        # «(-бесконечность;5)» и «( - бесконечность ; 5 )» — один ответ.
        self.assertEqual(len(summary["top_wrong"]), 1)
        self.assertEqual(summary["top_wrong"][0]["count"], 2)
        self.assertIn(summary["top_wrong"][0]["answer"], ("(-бесконечность;5)", "( - бесконечность ; 5 )"))

        status, data = self.review(task_id, "add_answer", answer="(-бесконечность; 5)")
        self.assertEqual(status, 200, data)
        self.assertEqual(data["extra_answers"], ["(-бесконечность; 5)"])
        self.assertEqual(data["attempts_changed"], 2)

        scores = self.scores()
        self.assertEqual(scores["Аня Первая"], (2, 2, 100))
        self.assertEqual(scores["Вера Третья"], (1, 2, 50))
        self.assertEqual(self.task_summary(task_id)["wrong"], 0)

        # Разбор работы: ответ теперь верный, среди правильных — добавленный.
        attempt_id = self.attempts["Аня Первая"]["attempt_id"]
        detail = self.teacher.get(f"/api/tests/{self.test_id}/attempts/{attempt_id}")[1]
        item = next(item for item in detail["items"] if item["task_id"] == task_id)
        self.assertTrue(item["is_correct"])
        self.assertIn("(-бесконечность; 5)", item["correct_answer"])
        self.assertEqual(item["extra_answers"], ["(-бесконечность; 5)"])

        # Умения и статистика пересчитаны.
        results = self.teacher.get(f"/api/tests/{self.test_id}/results")[1]
        skill = next(stat for stat in results["skill_stats"] if stat["position"] == 1)
        self.assertEqual((skill["correct"], skill["total"]), (2, 4))
        stats = self.teacher.get(f"/api/stats?test_id={self.test_id}")[1]
        self.assertEqual(stats["by_class"][0]["average_percent"], round((100 + 50 + 50 + 0) / 4))

        # Повтор — 409; повторный пересчёт ничего не меняет (идемпотентно).
        self.assertEqual(self.review(task_id, "add_answer", answer="(-БЕСКОНЕЧНОСТЬ;5)")[0], 409)
        from app import db
        from app.rescoring import rescore_test

        with db.get_pool().connection() as conn:
            self.assertEqual(rescore_test(conn, self.test_id), {"answers_changed": 0, "attempts_changed": 0})

        # Ответы учеников не тронуты.
        raw = query(
            "SELECT answer_text FROM answers WHERE attempt_id = %s AND task_id = %s", (attempt_id, task_id)
        )
        self.assertEqual(raw[0][0], "(-бесконечность;5)")

        # Отмена.
        status, data = self.review(task_id, "remove_answer", answer="(-бесконечность; 5)")
        self.assertEqual((status, data["extra_answers"]), (200, []))
        self.assertEqual(self.scores()["Аня Первая"], (1, 2, 50))
        # Исходный ответ задания удалить нельзя.
        self.assertEqual(self.review(task_id, "remove_answer", answer="бесконечность")[0], 404)

    def test_new_submissions_use_added_answer(self) -> None:
        task_id = self.attempts["Аня Первая"]["input_id"]
        self.review(task_id, "add_answer", answer="(-бесконечность;5)")
        late = self.submit("Дима Опоздавший", "(-бесконечность;5)", 0)
        if late["variant_no"] == 1:
            self.assertEqual((late["score"], late["max_score"]), (2, 2))

    def test_credit_all_and_exclude(self) -> None:
        input_task = self.attempts["Аня Первая"]["input_id"]

        self.assertEqual(self.review(input_task, "credit_all")[0], 200)
        scores = self.scores()
        self.assertEqual(scores["Аня Первая"], (2, 2, 100))
        self.assertEqual(scores["Вера Третья"], (1, 2, 50))  # выбор у Веры неверный
        # Второй вариант не затронут.
        self.assertEqual(scores["Гоша Четвёртый"], (1, 2, 50))
        self.assertEqual(self.review(input_task, "uncredit_all")[0], 200)
        self.assertEqual(self.scores()["Аня Первая"], (1, 2, 50))

        status, data = self.review(input_task, "exclude")
        self.assertEqual((status, data["excluded"]), (200, True))
        scores = self.scores()
        self.assertEqual(scores["Аня Первая"], (1, 1, 100))
        self.assertEqual(scores["Вера Третья"], (0, 1, 0))
        self.assertEqual(scores["Боря Второй"], (0, 2, 0))  # у варианта 2 задание своё

        results = self.teacher.get(f"/api/tests/{self.test_id}/results")[1]
        skill = next(stat for stat in results["skill_stats"] if stat["position"] == 1)
        self.assertEqual(skill["total"], 2)  # только вариант 2

        attempt_id = self.attempts["Аня Первая"]["attempt_id"]
        item = next(
            item
            for item in self.teacher.get(f"/api/tests/{self.test_id}/attempts/{attempt_id}")[1]["items"]
            if item["task_id"] == input_task
        )
        self.assertTrue(item["excluded"])

        # Excel: исключённое задание — «—», максимум меньше.
        status, content = self.teacher.get(f"/api/tests/{self.test_id}/export.xlsx", raw=True)
        rows = list(load_workbook(io.BytesIO(content))["Результаты"].iter_rows(values_only=True))
        anya = next(row for row in rows if row[1] == "Аня Первая")
        self.assertEqual(anya[3:], ("—", 1, 1, 1, 100))

        # Сдача после исключения тоже его не считает.
        late = self.submit("Женя Поздний", "что-то", 0)
        if late["variant_no"] == 1:
            self.assertEqual(late["max_score"], 1)
            self.assertTrue(next(r for r in late["results"] if r["task_id"] == input_task)["excluded"])

        self.assertEqual(self.review(input_task, "include")[0], 200)
        self.assertEqual(self.scores()["Аня Первая"], (1, 2, 50))

    def test_accept_option(self) -> None:
        choice_task = self.attempts["Вера Третья"]["choice_id"]
        wrong_option = self.attempts["Вера Третья"]["options"][1]["id"]
        right_option = self.attempts["Вера Третья"]["options"][0]["id"]

        summary = self.task_summary(choice_task)
        self.assertEqual(summary["top_wrong"], [{"answer": "6", "count": 1, "option_id": wrong_option, "students": ["Вера Т."], "more": 0}])

        self.assertEqual(self.review(choice_task, "accept_option", option_id=wrong_option)[0], 200)
        self.assertEqual(self.scores()["Вера Третья"], (1, 2, 50))
        summary = self.task_summary(choice_task)
        self.assertEqual(summary["wrong"], 0)
        self.assertTrue(next(o for o in summary["options"] if o["id"] == wrong_option)["accepted"])

        self.assertEqual(self.review(choice_task, "accept_option", option_id=right_option)[0], 400)
        self.assertEqual(self.review(choice_task, "add_answer", answer="6")[0], 400)
        self.assertEqual(self.review(choice_task, "unaccept_option", option_id=wrong_option)[0], 200)
        self.assertEqual(self.scores()["Вера Третья"], (0, 2, 0))

    def test_access_and_journal(self) -> None:
        task_id = self.attempts["Аня Первая"]["input_id"]
        stranger = Client().login("other-review@example.org")
        self.assertEqual(self.review(task_id, "exclude", client=stranger)[0], 403)
        self.assertEqual(stranger.get(f"/api/tests/{self.test_id}/tasks-review")[0], 403)
        self.assertEqual(stranger.get(f"/api/tests/{self.test_id}/review-log")[0], 403)

        admin = Client().login("review-admin@example.org")
        self.assertEqual(self.review(task_id, "exclude", client=admin)[0], 200)
        self.review(task_id, "include")

        # Задание другой работы по номеру — нет.
        other = self.teacher.post("/api/tests", body())[1]
        self.assertEqual(self.teacher.post(
            f"/api/tests/{other['id']}/tasks/{task_id}/review", {"action": "exclude"})[0], 404)
        self.assertEqual(self.review(task_id, "add_answer", answer="  ")[0], 422)
        self.assertEqual(self.review(task_id, "сломать")[0], 422)

        log = self.teacher.get(f"/api/tests/{self.test_id}/review-log")[1]
        self.assertEqual(
            [(row["user_name"], row["action"]) for row in log],
            [("Логарифмова Ольга", "include"), ("Главный Админ", "exclude")],
        )
        self.assertIn("вариант 1, задание 1", log[0]["detail"])
        self.assertEqual(log[1]["action_name"], "исключил задание из оценки")


class Migration011(unittest.TestCase):
    NAME = "kr_tests_tmp_migration011"

    def test_up_down_up(self) -> None:
        url = scratch_database(self.NAME)
        self.addCleanup(drop_database, self.NAME)
        up = next(path for path in migrations() if path.name.startswith("011"))
        down = up.with_name(up.name.replace(".sql", ".down.sql"))

        with psycopg.connect(url, autocommit=True) as conn:
            for path in migrations(up_to="010"):
                conn.execute(path.read_text(encoding="utf-8"))
            teacher = conn.execute(
                "INSERT INTO users (full_name, email, password_hash) VALUES ('У', 'u@e.org', 'x') RETURNING id"
            ).fetchone()[0]
            test_id = conn.execute(
                "INSERT INTO tests (title, teacher_id, teacher_name, share_token, is_published) "
                "VALUES ('Т', %s, 'У', 'abcdabcd', TRUE) RETURNING id",
                (teacher,),
            ).fetchone()[0]
            skill = conn.execute(
                "INSERT INTO skills (test_id, title, answer_format) VALUES (%s, 'У', 'choice') RETURNING id",
                (test_id,),
            ).fetchone()[0]
            task = conn.execute(
                "INSERT INTO tasks (test_id, skill_id, variant_no, text, answer_format) "
                "VALUES (%s, %s, 1, 'т', 'choice') RETURNING id",
                (test_id, skill),
            ).fetchone()[0]
            _, wrong = (
                conn.execute(
                    "INSERT INTO task_options (task_id, text, is_correct) VALUES (%s, %s, %s) RETURNING id",
                    (task, text, correct),
                ).fetchone()[0]
                for text, correct in (("да", True), ("нет", False))
            )
            attempt = conn.execute(
                "INSERT INTO attempts (test_id, student_name, attempt_token, finished_at, score, max_score) "
                "VALUES (%s, 'Ученик', 't', now(), 0, 1) RETURNING id",
                (test_id,),
            ).fetchone()[0]
            conn.execute(
                "INSERT INTO answers (attempt_id, task_id, option_id, is_correct) VALUES (%s, %s, %s, FALSE)",
                (attempt, task, wrong),
            )

            conn.execute(up.read_text(encoding="utf-8"))
            conn.execute(up.read_text(encoding="utf-8"))  # повторный накат безопасен
            # Учитель засчитал неверный вариант — как сделало бы приложение.
            conn.execute("UPDATE tasks SET extra_option_ids = ARRAY[%s]::bigint[] WHERE id = %s", (wrong, task))
            conn.execute("UPDATE answers SET is_correct = TRUE")
            conn.execute("UPDATE attempts SET score = 1")
            conn.execute(
                "INSERT INTO task_review_log (test_id, task_id, action) VALUES (%s, %s, 'accept_option')",
                (test_id, task),
            )

            conn.execute(down.read_text(encoding="utf-8"))
            columns = {
                row[0]
                for row in conn.execute("SELECT column_name FROM information_schema.columns WHERE table_name = 'tasks'")
            }
            self.assertFalse(columns & {"extra_answers", "extra_option_ids", "credit_all", "excluded"})
            # Вердикт и балл вернулись к исходному правильному ответу.
            self.assertEqual(conn.execute("SELECT is_correct FROM answers").fetchone()[0], False)
            self.assertEqual(conn.execute("SELECT score, max_score FROM attempts").fetchone(), (0, 1))
            self.assertIsNone(conn.execute("SELECT to_regclass('task_review_log')").fetchone()[0])

            conn.execute(up.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
